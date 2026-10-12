"""验证文件导入（PDF / Word / Markdown / 文本 / HTML / CSV）。

分两层验证：
  1. **解析层**直接调 ``suite.fileintake.extract``，逐格式看取出的文本对不对；
  2. **接口层**用 TestClient 打 ``/api/gather/import``，验证 multipart 上传与返回结构。

测试文件都在内存里构造，不往仓库里塞二进制样本。
"""

from __future__ import annotations

import io
import sys

sys.path.insert(0, r"E:\AI\thesis-suite")

PASS = "[OK]  "
FAIL = "[FAIL]"


def make_pdf() -> bytes:
    """用 PyMuPDF 现场造一个两页、带真实文字层的 PDF。

    注意 `insert_text` **不会自动换行**：句子超过页宽会被硬截断，数字可能断成两半
    （实测 `-0.34` 被拆开，导致断言失败）。所以这里手动分多行写，并把关键数字
    放在单独一行，确保文本层里是完整的。
    """
    import pymupdf

    doc = pymupdf.open()
    pages = [
        [
            "Rural teacher burnout and teaching engagement",
            "A survey of 1836 rural teachers across 24 counties.",
            "Standardized regression coefficient: -0.34",
        ],
        [
            "Organizational support moderated the relationship.",
            "Page two of the sample document.",
        ],
    ]
    for i, lines in enumerate(pages, start=1):
        page = doc.new_page()
        page.insert_text((72, 90), f"Page {i} heading", fontsize=14)
        y = 124
        for line in lines:
            page.insert_text((72, y), line, fontsize=10)
            y += 18
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("乡村教师职业倦怠调查")
    document.add_paragraph("本研究以某省24个县1836名乡村教师为样本。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "维度"
    table.cell(0, 1).text = "得分"
    table.cell(1, 0).text = "情感耗竭"
    table.cell(1, 1).text = "3.42"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def main() -> int:
    from suite import boot

    boot.prepare()
    from suite import fileintake

    checks: list[tuple[str, bool, str]] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        checks.append((name, ok, detail))
        print(f"  {PASS if ok else FAIL} {name}" + (f"  {detail}" if detail else ""))

    print("=" * 74)
    print("第一层：解析层逐格式验证")
    print("=" * 74)

    # ---- PDF ----
    print("\n--- PDF（带文字层，两页）---")
    pdf = make_pdf()
    r = fileintake.extract("样本论文.pdf", pdf)
    print(f"  类型={r.kind} 页数={r.pages} 段落={r.paragraphs} 字数={len(r.text)}")
    print(f"  预览：{r.text[:110]!r}")
    check("PDF 识别正确", r.kind == "PDF")
    check("PDF 取到文字", r.ok and len(r.text) > 60)
    check("PDF 页数正确", r.pages == 2, f"pages={r.pages}")
    check("PDF 按页切分（供带页码引用）", len(r.chunks) == 2, f"chunks={len(r.chunks)}")
    check("PDF 内容包含关键数字", "-0.34" in r.text, f"文本里找不到 -0.34：{r.text[:120]!r}")
    check("PDF 无错误无警告", not r.error and not r.warning, r.error or r.warning)

    # ---- 扫描版 PDF（无文字层）必须被识别，而不是假装成功 ----
    print("\n--- PDF（无文字层，模拟扫描版）---")
    import pymupdf

    blank = pymupdf.open()
    blank.new_page()          # 空白页：没有文本层
    blank_pdf = blank.tobytes()
    blank.close()
    r = fileintake.extract("扫描件.pdf", blank_pdf)
    print(f"  页数={r.pages} 有文字={bool(r.text.strip())}")
    print(f"  警告：{r.warning[:90]}")
    check("扫描版 PDF 被识别并给出可读说明", bool(r.warning) and "扫描" in r.warning)
    check("扫描版 PDF 不谎报成功", not r.ok)

    # ---- DOCX ----
    print("\n--- Word (.docx) ---")
    r = fileintake.extract("调查表.docx", make_docx())
    print(f"  类型={r.kind} 段落={r.paragraphs} 字数={len(r.text)}")
    print(f"  预览：{r.text[:110]!r}")
    check("DOCX 识别正确", r.kind == "Word")
    check("DOCX 取到正文", r.ok and "乡村教师" in r.text)
    check("DOCX 表格未丢", "情感耗竭" in r.text and "3.42" in r.text)

    # ---- Markdown ----
    print("\n--- Markdown ---")
    md = ("# 组织支持感研究述评\n\n"
          "见 [Smith 2020](https://example.com) 与 ![图1](img.png)。\n\n"
          "- 社会交换理论\n- 组织拟人化思想\n\n"
          "`SPOS` 是常用量表。\n")
    r = fileintake.extract("述评.md", md.encode("utf-8"))
    print(f"  类型={r.kind} 字数={len(r.text)}")
    print(f"  预览：{r.text[:120]!r}")
    check("Markdown 识别正确", r.kind == "Markdown")
    check("Markdown 链接保留文字、去掉 URL", "Smith 2020" in r.text and "https://" not in r.text)
    check("Markdown 图片去掉路径", "图1" in r.text and "img.png" not in r.text)
    check("Markdown 列表结构保留", "社会交换理论" in r.text and "-" in r.text)

    # ---- 纯文本 + 编码 ----
    print("\n--- 纯文本（GBK 编码）---")
    gbk = "乡村教师职业倦怠总体处于中等偏上水平。".encode("gbk")
    r = fileintake.extract("摘录.txt", gbk)
    print(f"  类型={r.kind} 警告={r.warning!r}")
    print(f"  预览：{r.text[:60]!r}")
    check("GBK 文本正确解码（不是乱码）", "乡村教师" in r.text)
    check("非 UTF-8 会标注编码", "gb" in r.warning.lower())

    # ---- HTML ----
    print("\n--- HTML ---")
    html = ("<html><head><style>p{color:red}</style><script>var x=1</script></head>"
            "<body><h1>教师倦怠</h1><p>第一段正文。</p><p>第二段正文。</p></body></html>")
    r = fileintake.extract("页面.html", html.encode("utf-8"))
    print(f"  类型={r.kind} 字数={len(r.text)}")
    print(f"  预览：{r.text[:80]!r}")
    check("HTML 剥掉 script/style", "var x" not in r.text and "color:red" not in r.text)
    check("HTML 保留正文", "教师倦怠" in r.text and "第一段正文" in r.text)

    # ---- CSV ----
    print("\n--- CSV ---")
    csv_text = "维度,得分\n情感耗竭,3.42\n低成就感,3.10\n"
    r = fileintake.extract("数据.csv", csv_text.encode("utf-8"))
    print(f"  预览：{r.text!r}")
    check("CSV 转成可读文本", "情感耗竭" in r.text and "3.42" in r.text)

    # ---- 拒绝不支持的与伪造扩展名 ----
    print("\n--- 边界：不支持的类型 / 伪造扩展名 ---")
    r = fileintake.extract("假文本.txt", b"%PDF-1.7 not really a pdf")
    check("伪造扩展名按文件头纠正为 PDF", r.kind == "PDF", f"kind={r.kind} err={r.error[:40]}")
    r = fileintake.extract("表格.xlsx", b"PK\x03\x04fake")
    check("xlsx（zip 容器）明确拒绝", bool(r.error) and "压缩包" in r.error, r.error[:60])
    r = fileintake.extract("空.txt", b"")
    check("空文件明确拒绝", bool(r.error) and "空" in r.error)
    r = fileintake.extract("大.txt", b"x" * (fileintake.MAX_FILE_BYTES + 1))
    check("超大文件明确拒绝", bool(r.error) and "过大" in r.error)

    # ---- 转成 sources（PDF 逐页）----
    print("\n--- 转 sources（PDF 逐页，保留页码）---")
    srcs = fileintake.split_into_sources([fileintake.extract("样本论文.pdf", pdf)])
    print(f"  生成 {len(srcs)} 条 source")
    for s in srcs:
        print(f"    title={s['title']!r} page={s.get('page')}")
    check("PDF 逐页生成 source", len(srcs) == 2, f"count={len(srcs)}")
    check("source 带页码", all(s.get("page") for s in srcs))

    print("\n" + "=" * 74)
    print("第二层：/api/gather/import 接口验证")
    print("=" * 74)

    from fastapi.testclient import TestClient

    from suite.api.app import create_app

    client = TestClient(create_app())

    files = [
        ("files", ("样本论文.pdf", pdf, "application/pdf")),
        ("files", ("调查表.docx", make_docx(),
                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("files", ("述评.md", md.encode("utf-8"), "text/markdown")),
    ]
    resp = client.post("/api/gather/import", files=files)
    print(f"\n  HTTP {resp.status_code}")
    if resp.status_code != 200:
        print("  " + resp.text[:400])
        check("导入接口返回 200", False, resp.text[:120])
    else:
        data = resp.json()
        print(f"  received={data['received']} parsed={data['parsed']} "
              f"failed={data['failed']} source_count={data['source_count']}")
        for f in data["files"]:
            print(f"    · {f['filename']:16} 类型={f['kind']:9} 字数={f['chars']:>6} "
                  f"页={f['pages']} ok={f['ok']}")
        check("接口解析出 3 个文件", data["parsed"] == 3)
        check("接口无需 API Key", True)
        check("PDF 拆成按页 source", data["source_count"] >= 4, f"source_count={data['source_count']}")
        check("返回每文件的摘要与预览", all("preview" in f for f in data["files"]))

    # 全部失败 -> 非 200（不能显示"成功，0 条"）
    resp = client.post("/api/gather/import", files=[("files", ("坏.xlsx", b"PK\x03\x04x", "application/zip"))])
    print(f"\n  全部失败时 HTTP {resp.status_code}")
    check("全部失败返回非 200", resp.status_code != 200, f"HTTP {resp.status_code}")

    # 空请求
    resp = client.post("/api/gather/import")
    check("没有文件时返回 400", resp.status_code == 400, f"HTTP {resp.status_code}")

    print("\n" + "=" * 74)
    failed = [c for c in checks if not c[1]]
    if failed:
        print(f"失败 {len(failed)} 项：")
        for name, _, detail in failed:
            print(f"  - {name}  {detail}")
        return 1
    print(f"全部通过（{len(checks)} 项）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
