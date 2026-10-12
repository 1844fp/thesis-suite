"""通过通用 agent 接口验证本轮补齐的五个功能。

**不调用模型**（不需要花 API Key 的钱）：
  * Word 导出、查重自查、模板比对 —— 本来就不需要 Key
  * 大纲生成与按意见修改 —— 需要 Key，这里只验证"缺 Key 时给出可读提示"
    （真模型的验证由 tools/smoke_keyonly.py 之类的脚本负责）
"""

from __future__ import annotations

import base64
import io
import sys

sys.path.insert(0, r"E:\AI\thesis-suite")

ok = True


def check(name: str, cond: bool, detail: str = "") -> None:
    global ok
    print(f"  [{'OK ' if cond else 'FAIL'}] {name}" + (f"  {detail}" if detail else ""))
    ok = ok and cond


def make_docx(east_font: str, size: float, spacing) -> bytes:
    import docx
    from docx.oxml.ns import qn

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.size = docx.shared.Pt(size)
    st.paragraph_format.line_spacing = spacing
    rpr = st.element.get_or_add_rPr()
    rf = rpr.makeelement(qn("w:rFonts"), {})
    rpr.append(rf)
    rf.set(qn("w:eastAsia"), east_font)
    d.add_paragraph("这是一段用于模板识别的正文内容，需要足够长才会被当作正文。")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def main() -> int:
    from fastapi.testclient import TestClient

    from suite.api.app import create_app

    client = TestClient(create_app())

    print("=" * 70)
    print("① 大纲生成（需要 Key，这里验证缺 Key 时的提示与参数校验）")
    print("=" * 70)
    r = client.post("/api/agents/topic/outline", json={"title": "乡村教师职业倦怠研究"})
    check("缺 Key 时拒绝并提示可读", r.status_code == 400 and "API Key" in r.text,
          f"HTTP {r.status_code} {r.text[:80]}")
    r = client.post("/api/agents/topic/outline", json={"title": ""},
                    headers={"X-LLM-Key": "sk-fake"})
    check("缺题目时提示先给题目", r.status_code == 200 and r.json().get("ok") is False,
          str(r.json().get("error"))[:60])

    print("\n" + "=" * 70)
    print("② Word 导出（不需要 Key）")
    print("=" * 70)
    # 先造一个会话与稿件：借用 topic/generate 建会话，再直接写版本
    r = client.post("/api/agents/topic/generate",
                    json={"discipline": "教育学", "count": 1, "mode": "offline"})
    sid = r.json().get("sid")
    check("建会话成功", bool(sid))

    from suite.sessions import STORE

    sess = STORE.get(sid)
    sess.versions.append({
        "title": "乡村教师职业倦怠研究",
        "sections": [
            {"heading": "第1章 绪论", "body": "乡村教师队伍建设是乡村教育质量的基础环节[1]。"},
            {"heading": "第2章 文献综述", "body": "现有研究结论并不一致[2]。"},
        ],
        "words": 60,
        "coverage": {"sections": 2, "citations_total": 2, "citations_user": 2,
                     "citations_web": 0, "citations_invalid": 0, "ungrounded_numbers": 0,
                     "verdict": "仅依据你提供的资料"},
    })

    r = client.post("/api/agents/writer/export_docx", json={"sid": sid})
    check("导出 Word 200", r.status_code == 200, f"HTTP {r.status_code}")
    data = r.json()
    check("导出成功", data.get("ok") is True, str(data.get("error"))[:80])
    raw = base64.b64decode(data.get("content_base64") or "")
    check("是合法 docx（zip 头）", raw[:2] == b"PK", f"{len(raw)} 字节")
    check("文件名以 .docx 结尾", str(data.get("filename", "")).endswith(".docx"),
          str(data.get("filename")))

    # 反向读回，确认内容真的写进去了
    import docx

    d = docx.Document(io.BytesIO(raw))
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    check("Word 里有标题", any("乡村教师职业倦怠研究" in t for t in texts), str(texts[:1]))
    check("Word 里有章节标题", any("第1章" in t for t in texts))
    heads = [p.style.name for p in d.paragraphs if p.style.name.startswith("Heading")]
    check("章节用了 Heading 样式（可被模板替换）", bool(heads), str(heads[:3]))

    print("\n" + "=" * 70)
    print("③ 查重风险自查（不需要 Key）")
    print("=" * 70)
    dup_text = (
        "乡村教师职业倦怠已成为影响教学质量的关键因素，本研究对此展开系统考察并给出政策建议。\n"
        "组织支持感能够缓冲工作压力对工作投入的负面影响，这一点在既有研究中得到一致支持。\n"
        "乡村教师职业倦怠已成为影响教学质量的关键因素，本研究对此展开系统考察并给出政策建议。"
    )
    r = client.post("/api/agents/review/duplication", json={"sid": sid, "text": dup_text})
    check("查重自查 200", r.status_code == 200, f"HTTP {r.status_code}")
    data = r.json()
    check("查重成功", data.get("ok") is True, str(data.get("error"))[:80])
    internal = data.get("internal") or {}
    check("检出内部重复", internal.get("duplicate_pair_count", 0) >= 1,
          f"段落={internal.get('paragraphs')} 重复对={internal.get('duplicate_pair_count')}")
    check("风险等级不是 low", data.get("risk") in ("medium", "high"), str(data.get("risk")))
    check("带上了免责声明（不能替代学校查重）",
          "不是学校查重" in str(data.get("disclaimer")),
          str(data.get("disclaimer"))[:60])

    print("\n" + "=" * 70)
    print("④ 学校模板比对（不需要 Key）")
    print("=" * 70)
    tpl = make_docx("宋体", 12, 1.5)
    man = make_docx("楷体", 10.5, 1.0)
    r = client.post("/api/agents/review/template", json={
        "sid": sid,
        "template_b64": base64.b64encode(tpl).decode("ascii"),
        "manuscript_b64": base64.b64encode(man).decode("ascii"),
    })
    check("模板比对 200", r.status_code == 200, f"HTTP {r.status_code}")
    data = r.json()
    check("比对成功", data.get("ok") is True, str(data.get("error"))[:80])
    check("识别出模板字号", (data.get("template") or {}).get("styles", {}).get("Normal", {}).get("size_pt") == 12.0,
          str((data.get("template") or {}).get("styles")))
    check("找出不合规项", len(data.get("diffs") or []) >= 2, f"{len(data.get('diffs') or [])} 处")
    check("判定不合规", data.get("conform") is False)

    print("\n" + "=" * 70)
    print("⑤ 按意见改写（需要 Key，验证缺 Key 与缺稿件时的提示）")
    print("=" * 70)
    r = client.post("/api/agents/writer/revise", json={"sid": sid, "heading": "第1章"})
    check("缺 Key 时拒绝并提示", r.status_code == 400 and "API Key" in r.text,
          f"HTTP {r.status_code}")
    # 给假 Key 时应该走到"章节匹配"那一步
    r = client.post("/api/agents/writer/revise", json={"sid": sid, "heading": "不存在的章"},
                    headers={"X-LLM-Key": "sk-fake"})
    body = r.json()
    check("章节不存在时列出可用章节", body.get("ok") is False and "可用的章节" in str(body.get("error")),
          str(body.get("error"))[:70])

    print("\n" + "=" * 70)
    print("⑥ 事件流应记录下以上所有动作")
    print("=" * 70)
    r = client.get(f"/api/agents/events?sid={sid}")
    events = r.json().get("events") or []
    agents_seen = {e["agent"] for e in events}
    check("事件流记录了多个 agent", len(agents_seen) >= 2, str(sorted(agents_seen)))
    actions = {e["action"] for e in events}
    check("记录了新功能的动作", {"export_docx", "duplication", "template"} <= actions,
          str(sorted(actions)))

    print("\n" + "=" * 70)
    print("全部通过" if ok else "有失败项")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
